#!/usr/bin/env python3
"""Convert WORKING_INSTRUCTION.md to WORKING_INSTRUCTION.docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "WORKING_INSTRUCTION.md"
DOCX_PATH = ROOT / "WORKING_INSTRUCTION.docx"


def add_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def main() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md.splitlines()
    i = 0
    in_table = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows = [
            r
            for r in table_rows
            if not all(re.fullmatch(r":?-{3,}:?", c.strip() or "") for c in r)
        ]
        table_rows = []
        in_table = False
        if not rows:
            return
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell_text = row[ci].strip() if ci < len(row) else ""
                cell = table.rows[ri].cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs(p, cell_text)
                if ri == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            i += 1
            continue
        if in_table:
            flush_table()

        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("---"):
            pass
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif line.startswith("  - ") or line.startswith("   - "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, line.lstrip()[2:])
        elif line.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
        i += 1

    if in_table:
        flush_table()

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
